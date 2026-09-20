import pytest
import uuid
import io
import openpyxl
from datetime import datetime, timezone, timedelta
from sqlalchemy import select
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker, AsyncSession
from sqlalchemy.pool import NullPool

from backend.app.core.config import settings
from backend.app.models.user import User, UserRole
from backend.app.models.question import Question, QuestionDifficulty
from backend.app.models.exam import Exam, ExamAssignment, AssignedQuestion, AssignmentStatus
from backend.app.models.submission import Submission, MCQResponse
from backend.app.models.proctoring import ExamProctoringLog
from backend.app.models.result import ExamResult
from backend.app.services.exam_service import (
    get_exam_leaderboard_paginated,
    export_exam_leaderboard_excel,
)


@pytest.mark.asyncio
async def test_leaderboard_pagination_and_sorting():
    """Verify server-side pagination math and sorting."""
    test_engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
    test_session = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with test_session() as db:
        now = datetime.now(timezone.utc)
        suffix = uuid.uuid4().hex[:6]

        # 1. Create Exam
        exam = Exam(
            title=f"Pagination Test Exam {suffix}",
            duration_minutes=60,
            easy_count=1,
            medium_count=1,
            hard_count=0,
            mcq_count=1,
        )
        db.add(exam)
        await db.flush()

        # 2. Seed 5 students with assignments and varying scores/ranks
        students = []
        for i in range(5):
            student = User(
                name=f"Student {i+1} {suffix}",
                email=f"student_{i+1}_{suffix}@test.com",
                roll_no=f"ROLL_{i+1}_{suffix}",
                role=UserRole.STUDENT,
                college="Tech University",
                candidate_group="Batch-2026",
                password_hash="pw",
            )
            db.add(student)
            await db.flush()
            students.append(student)

            assign = ExamAssignment(
                exam_id=exam.id,
                user_id=student.id,
                status=AssignmentStatus.SUBMITTED,
                started_at=now - timedelta(minutes=45),
                submitted_at=now - timedelta(minutes=45 - (i * 5)),
                is_active=True,
            )
            db.add(assign)
            await db.flush()

            # Assign results with rank and score
            result = ExamResult(
                assignment_id=assign.id,
                total_score=float(100 - (i * 15)),  # 100, 85, 70, 55, 40
                rank=i + 1,
            )
            db.add(result)
            await db.flush()

        await db.commit()

        # 3. Test page 1, page_size 2
        page1 = await get_exam_leaderboard_paginated(
            db=db, exam_id=exam.id, page=1, page_size=2, sort_by="rank", sort_dir="asc"
        )
        assert page1.total_count == 5
        assert page1.total_pages == 3
        assert page1.current_page == 1
        assert page1.page_size == 2
        assert len(page1.items) == 2
        assert page1.items[0].rank == 1
        assert page1.items[1].rank == 2

        # 4. Test page 2, page_size 2
        page2 = await get_exam_leaderboard_paginated(
            db=db, exam_id=exam.id, page=2, page_size=2, sort_by="rank", sort_dir="asc"
        )
        assert page2.current_page == 2
        assert len(page2.items) == 2
        assert page2.items[0].rank == 3
        assert page2.items[1].rank == 4

        # 5. Test page 3, page_size 2
        page3 = await get_exam_leaderboard_paginated(
            db=db, exam_id=exam.id, page=3, page_size=2, sort_by="rank", sort_dir="asc"
        )
        assert page3.current_page == 3
        assert len(page3.items) == 1
        assert page3.items[0].rank == 5

        # 6. Test search filter
        search_res = await get_exam_leaderboard_paginated(
            db=db, exam_id=exam.id, page=1, page_size=10, search=f"Student 3 {suffix}"
        )
        assert search_res.total_count == 1
        assert search_res.items[0].name == f"Student 3 {suffix}"


@pytest.mark.asyncio
async def test_violation_aggregation():
    """Verify proctoring logs are aggregated correctly by type in a single grouped query."""
    test_engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
    test_session = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with test_session() as db:
        now = datetime.now(timezone.utc)
        suffix = uuid.uuid4().hex[:6]

        exam = Exam(title=f"Violation Exam {suffix}", duration_minutes=60)
        db.add(exam)
        await db.flush()

        student = User(
            name=f"Cheater Candidate {suffix}",
            email=f"cheat_{suffix}@test.com",
            roll_no=f"CHT_{suffix}",
            role=UserRole.STUDENT,
            password_hash="pw",
        )
        db.add(student)
        await db.flush()

        assign = ExamAssignment(
            exam_id=exam.id,
            user_id=student.id,
            status=AssignmentStatus.IN_PROGRESS,
            started_at=now,
            is_active=True,
        )
        db.add(assign)
        await db.flush()

        # Add proctoring logs (including benign device logs that should NOT count as infractions)
        logs = [
            ExamProctoringLog(assignment_id=assign.id, event_type="TAB_SWITCH", title="Tab", description="Tab switch"),
            ExamProctoringLog(assignment_id=assign.id, event_type="TAB_SWITCH", title="Tab", description="Tab switch 2"),
            ExamProctoringLog(assignment_id=assign.id, event_type="FULLSCREEN_EXIT", title="FS", description="Exit FS"),
            ExamProctoringLog(assignment_id=assign.id, event_type="WINDOW_BLUR", title="Blur", description="Lost focus"),
            ExamProctoringLog(assignment_id=assign.id, event_type="PASTE_ATTEMPT", title="Paste", description="Pasted code"),
            ExamProctoringLog(assignment_id=assign.id, event_type="DEVTOOLS_SHORTCUT", title="DevTools", description="Inspected"),
            ExamProctoringLog(assignment_id=assign.id, event_type="NAVIGATION_BLOCKED", title="Nav", description="Back attempt"),
            # Informational / non-infraction event types that must NOT increment total_violation_count:
            ExamProctoringLog(assignment_id=assign.id, event_type="EXAM_START_DEVICE", title="Start", description="Initial camera"),
            ExamProctoringLog(assignment_id=assign.id, event_type="VERIFICATION_SNAPSHOT", title="Snap", description="Snap"),
        ]
        db.add_all(logs)
        await db.commit()

        res = await get_exam_leaderboard_paginated(db=db, exam_id=exam.id, page=1, page_size=10)
        assert res.total_count == 1
        entry = res.items[0]

        assert entry.total_violation_count == 7  # 2 tab + 1 fs + 1 blur + 1 paste + 1 devtools + 1 nav
        assert entry.tab_switch_count == 2
        assert entry.fullscreen_exit_count == 1
        assert entry.blur_count == 1
        assert entry.clipboard_block_count == 1
        assert entry.devtools_attempt_count == 1
        assert entry.navigation_block_count == 1


@pytest.mark.asyncio
async def test_coding_solved_and_mcq_correct_breakdown():
    """Verify coding-solved count (correctness==1.0 on final) and MCQ correct count."""
    test_engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
    test_session = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with test_session() as db:
        now = datetime.now(timezone.utc)
        suffix = uuid.uuid4().hex[:6]

        exam = Exam(
            title=f"Breakdown Exam {suffix}",
            duration_minutes=60,
            easy_count=2,
            medium_count=0,
            hard_count=0,
            mcq_count=2,
        )
        db.add(exam)
        await db.flush()

        # Create 2 coding questions
        cq1 = Question(title=f"CQ1 {suffix}", description="CQ1", difficulty=QuestionDifficulty.EASY, question_type="coding")
        cq2 = Question(title=f"CQ2 {suffix}", description="CQ2", difficulty=QuestionDifficulty.EASY, question_type="coding")
        # Create 2 MCQ questions
        mq1 = Question(title=f"MQ1 {suffix}", description="MQ1", difficulty=QuestionDifficulty.EASY, question_type="mcq", marks=5.0)
        mq2 = Question(title=f"MQ2 {suffix}", description="MQ2", difficulty=QuestionDifficulty.EASY, question_type="mcq", marks=5.0)
        db.add_all([cq1, cq2, mq1, mq2])
        await db.flush()

        student = User(
            name=f"Solver Candidate {suffix}",
            email=f"solver_{suffix}@test.com",
            role=UserRole.STUDENT,
            password_hash="pw",
        )
        db.add(student)
        await db.flush()

        assign = ExamAssignment(
            exam_id=exam.id,
            user_id=student.id,
            status=AssignmentStatus.SUBMITTED,
            started_at=now - timedelta(minutes=30),
            submitted_at=now,
            is_active=True,
        )
        db.add(assign)
        await db.flush()

        # Assigned questions
        aq_list = [
            AssignedQuestion(assignment_id=assign.id, question_id=cq1.id, difficulty=QuestionDifficulty.EASY),
            AssignedQuestion(assignment_id=assign.id, question_id=cq2.id, difficulty=QuestionDifficulty.EASY),
            AssignedQuestion(assignment_id=assign.id, question_id=mq1.id, difficulty=QuestionDifficulty.EASY),
            AssignedQuestion(assignment_id=assign.id, question_id=mq2.id, difficulty=QuestionDifficulty.EASY),
        ]
        db.add_all(aq_list)

        # Submission for CQ1: 4/4 passed (SOLVED)
        s1 = Submission(
            assignment_id=assign.id,
            question_id=cq1.id,
            code="print(1)",
            language="python",
            status="Accepted",
            test_cases_passed=4,
            total_test_cases=4,
            is_final=True,
        )
        # Submission for CQ2: 2/4 passed (NOT solved)
        s2 = Submission(
            assignment_id=assign.id,
            question_id=cq2.id,
            code="print(2)",
            language="python",
            status="Wrong Answer",
            test_cases_passed=2,
            total_test_cases=4,
            is_final=True,
        )
        # MCQ response for MQ1: Correct
        mr1 = MCQResponse(
            assignment_id=assign.id,
            question_id=mq1.id,
            selected_option_ids=[],
            is_correct=True,
            marks_awarded=5.0,
        )
        # MCQ response for MQ2: Incorrect
        mr2 = MCQResponse(
            assignment_id=assign.id,
            question_id=mq2.id,
            selected_option_ids=[],
            is_correct=False,
            marks_awarded=0.0,
        )
        db.add_all([s1, s2, mr1, mr2])
        await db.commit()

        res = await get_exam_leaderboard_paginated(db=db, exam_id=exam.id, page=1, page_size=10)
        assert res.total_count == 1
        row = res.items[0]

        assert row.questions_solved_count == 1
        assert row.total_coding_questions == 2
        assert row.mcq_correct_count == 1
        assert row.total_mcq_questions == 2


@pytest.mark.asyncio
async def test_not_started_candidate_null_score_and_rank():
    """Verify candidates with status='not_started' appear with null score/rank and don't crash the query."""
    test_engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
    test_session = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with test_session() as db:
        now = datetime.now(timezone.utc)
        suffix = uuid.uuid4().hex[:6]

        exam = Exam(title=f"Unstarted Exam {suffix}", duration_minutes=45)
        db.add(exam)
        await db.flush()

        student = User(
            name=f"Absent Student {suffix}",
            email=f"absent_{suffix}@test.com",
            role=UserRole.STUDENT,
            password_hash="pw",
        )
        db.add(student)
        await db.flush()

        assign = ExamAssignment(
            exam_id=exam.id,
            user_id=student.id,
            status=AssignmentStatus.NOT_STARTED,
            is_active=True,
        )
        db.add(assign)
        await db.commit()

        res = await get_exam_leaderboard_paginated(db=db, exam_id=exam.id, page=1, page_size=10)
        assert res.total_count == 1
        row = res.items[0]

        assert row.status == "not_started"
        assert row.total_score is None
        assert row.rank is None
        assert row.started_at is None
        assert row.submitted_at is None
        assert row.time_taken_seconds is None


@pytest.mark.asyncio
async def test_export_excel_structure_and_filters():
    """Verify Excel export produces a valid .xlsx file with proper headers, styling, and data."""
    test_engine = create_async_engine(settings.DATABASE_URL, poolclass=NullPool)
    test_session = async_sessionmaker(bind=test_engine, class_=AsyncSession, expire_on_commit=False)

    async with test_session() as db:
        now = datetime.now(timezone.utc)
        suffix = uuid.uuid4().hex[:6]

        exam = Exam(title=f"Export Assessment {suffix}", duration_minutes=60)
        db.add(exam)
        await db.flush()

        # Seed 2 students: one submitted, one in_progress
        u1 = User(
            name=f"Export Candidate 1 {suffix}",
            email=f"exp1_{suffix}@test.com",
            roll_no=f"ROLL1_{suffix}",
            college="MIT",
            candidate_group="Batch-A",
            role=UserRole.STUDENT,
            password_hash="pw",
        )
        u2 = User(
            name=f"Export Candidate 2 {suffix}",
            email=f"exp2_{suffix}@test.com",
            roll_no=f"ROLL2_{suffix}",
            college="Stanford",
            candidate_group="Batch-B",
            role=UserRole.STUDENT,
            password_hash="pw",
        )
        db.add_all([u1, u2])
        await db.flush()

        a1 = ExamAssignment(
            exam_id=exam.id,
            user_id=u1.id,
            status=AssignmentStatus.SUBMITTED,
            started_at=now - timedelta(minutes=40),
            submitted_at=now - timedelta(minutes=10),
            is_active=True,
        )
        a2 = ExamAssignment(
            exam_id=exam.id,
            user_id=u2.id,
            status=AssignmentStatus.IN_PROGRESS,
            started_at=now - timedelta(minutes=20),
            is_active=True,
        )
        db.add_all([a1, a2])
        await db.flush()

        # Result for a1
        res1 = ExamResult(assignment_id=a1.id, total_score=95.0, rank=1)
        db.add(res1)
        await db.commit()

        # Helper to read StreamingResponse bytes from async_generator
        async def read_stream_bytes(s):
            chunks = []
            async for chunk in s.body_iterator:
                chunks.append(chunk)
            return b"".join(chunks)

        # 1. Unfiltered Export: 2 rows + 1 header = 3 rows
        stream = await export_exam_leaderboard_excel(db=db, exam_id=exam.id)
        content = await read_stream_bytes(stream)

        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active

        assert ws.title == "Leaderboard Results"
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None

        # Verify Header Columns
        header_row = [cell.value for cell in ws[1]]
        expected_cols = [
            "Rank", "Name", "Email", "Roll No", "College", "Group", "Status",
            "Total Score", "Coding Solved", "MCQ Correct", "Time Taken", "Started At",
            "Submitted At", "Total Violations", "Tab Switches", "Fullscreen Exits",
            "Window Blurs", "Clipboard Blocks", "DevTools Attempts", "Navigation Blocks"
        ]
        assert header_row == expected_cols

        # Data rows count
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 2

        # 2. Filtered Export (status="submitted")
        stream_filtered = await export_exam_leaderboard_excel(db=db, exam_id=exam.id, status="submitted")
        content_filtered = await read_stream_bytes(stream_filtered)

        wb_filtered = openpyxl.load_workbook(io.BytesIO(content_filtered))
        ws_filtered = wb_filtered.active
        filtered_rows = list(ws_filtered.iter_rows(min_row=2, values_only=True))
        assert len(filtered_rows) == 1
        assert filtered_rows[0][1] == f"Export Candidate 1 {suffix}"
